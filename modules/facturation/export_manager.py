"""
Export Manager for PMO Billing System
Modern export interface with preview, scheduling, and multiple formats
"""

import streamlit as st
import pandas as pd
import numpy as np
import json
import csv
import io
import base64
from datetime import datetime, date, timedelta
from typing import Dict, List, Optional, Any, Union, Tuple
import zipfile
import tempfile
import os
from pathlib import Path

# PDF generation
try:
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib import colors
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# Excel generation
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils.dataframe import dataframe_to_rows
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False


class ExportManager:
    """Modern export manager with multiple formats and preview capabilities"""
    
    def __init__(self):
        self.supported_formats = self._get_supported_formats()
        self.init_session_state()
    
    def init_session_state(self):
        """Initialize session state for export management"""
        if 'export_history' not in st.session_state:
            st.session_state.export_history = []
        if 'scheduled_exports' not in st.session_state:
            st.session_state.scheduled_exports = []
        if 'export_templates' not in st.session_state:
            st.session_state.export_templates = self._get_default_templates()
    
    def _get_supported_formats(self) -> Dict[str, Dict[str, Any]]:
        """Get list of supported export formats"""
        formats = {
            "CSV": {
                "description": "Comma Separated Values - Compatible avec Excel et autres tableurs",
                "icon": "📊",
                "available": True,
                "mime_type": "text/csv",
                "extension": ".csv"
            },
            "JSON": {
                "description": "JavaScript Object Notation - Format structuré pour les APIs",
                "icon": "📋",
                "available": True,
                "mime_type": "application/json",
                "extension": ".json"
            },
            "Excel": {
                "description": "Microsoft Excel - Feuille de calcul avec mise en forme",
                "icon": "📈",
                "available": EXCEL_AVAILABLE,
                "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "extension": ".xlsx"
            },
            "PDF": {
                "description": "Portable Document Format - Document prêt à imprimer",
                "icon": "📄",
                "available": PDF_AVAILABLE,
                "mime_type": "application/pdf",
                "extension": ".pdf"
            },
            "HTML": {
                "description": "HyperText Markup Language - Page web avec styles",
                "icon": "🌐",
                "available": True,
                "mime_type": "text/html",
                "extension": ".html"
            }
        }
        
        return formats
    
    def _get_default_templates(self) -> Dict[str, Dict[str, Any]]:
        """Get default export templates"""
        return {
            "Factures Mensuelles": {
                "description": "Export des factures pour une période donnée",
                "columns": ["participant_name", "period", "amount_ht", "amount_ttc", "status"],
                "filters": {"period": "monthly"},
                "format": "Excel"
            },
            "Rapport Analytique": {
                "description": "Analyse complète avec graphiques et KPIs",
                "columns": ["all"],
                "filters": {},
                "format": "PDF"
            },
            "Données Brutes": {
                "description": "Export complet de toutes les données",
                "columns": ["all"],
                "filters": {},
                "format": "CSV"
            },
            "Résumé Exécutif": {
                "description": "Résumé pour la direction avec métriques clés",
                "columns": ["summary"],
                "filters": {"level": "executive"},
                "format": "PDF"
            }
        }
    
    def create_export_interface(self, data: pd.DataFrame, data_type: str = "generic"):
        """Create modern export interface"""
        
        st.markdown("### 📤 Centre d'Export Moderne")
        st.markdown("*Exportez vos données dans le format de votre choix avec aperçu en temps réel*")
        
        if data.empty:
            st.warning("⚠️ Aucune donnée disponible pour l'export")
            return
        
        # Export tabs
        tab1, tab2, tab3, tab4 = st.tabs([
            "🚀 Export Rapide", 
            "⚙️ Export Personnalisé", 
            "📅 Export Programmé", 
            "📊 Historique"
        ])
        
        with tab1:
            self._create_quick_export(data, data_type)
        
        with tab2:
            self._create_custom_export(data, data_type)
        
        with tab3:
            self._create_scheduled_export(data, data_type)
        
        with tab4:
            self._show_export_history()
    
    def _create_quick_export(self, data: pd.DataFrame, data_type: str):
        """Create quick export interface"""
        
        st.markdown("#### ⚡ Export Rapide")
        st.markdown("Exportez rapidement vos données avec les paramètres par défaut")
        
        # Format selection with visual cards
        st.markdown("##### Choisir le Format")
        
        format_cols = st.columns(len([f for f in self.supported_formats.values() if f['available']]))
        selected_format = None
        
        for i, (format_name, format_info) in enumerate([
            (name, info) for name, info in self.supported_formats.items() 
            if info['available']
        ]):
            with format_cols[i]:
                if st.button(
                    f"{format_info['icon']}\n**{format_name}**\n{format_info['description'][:30]}...",
                    key=f"quick_export_{format_name}",
                    use_container_width=True
                ):
                    selected_format = format_name
        
        if selected_format:
            self._process_export(data, selected_format, {}, f"export_rapide_{data_type}")
    
    def _create_custom_export(self, data: pd.DataFrame, data_type: str):
        """Create custom export interface"""
        
        st.markdown("#### 🛠️ Export Personnalisé")
        
        col1, col2 = st.columns([2, 1])
        
        with col1:
            # Template selection
            template_name = st.selectbox(
                "📋 Modèle d'Export",
                ["Personnalisé"] + list(st.session_state.export_templates.keys()),
                help="Sélectionnez un modèle prédéfini ou créez un export personnalisé"
            )
            
            # Load template if selected
            template_config = {}
            if template_name != "Personnalisé":
                template_config = st.session_state.export_templates[template_name]
                st.info(f"📄 **{template_name}:** {template_config['description']}")
        
        with col2:
            # Save as template option
            if st.button("💾 Sauvegarder comme Modèle"):
                self._save_export_template()
        
        # Format selection
        format_name = st.selectbox(
            "📁 Format d'Export",
            [name for name, info in self.supported_formats.items() if info['available']],
            index=0 if template_config.get('format') not in self.supported_formats 
                  else list(self.supported_formats.keys()).index(template_config.get('format', 'CSV')),
            help="Sélectionnez le format de fichier désiré"
        )
        
        # Column selection
        st.markdown("##### 📋 Sélection des Colonnes")
        
        available_columns = list(data.columns)
        
        if template_config.get('columns') == ['all']:
            default_columns = available_columns
        elif template_config.get('columns'):
            default_columns = [col for col in template_config['columns'] if col in available_columns]
        else:
            default_columns = available_columns[:5]  # First 5 columns by default
        
        selected_columns = st.multiselect(
            "Colonnes à exporter",
            available_columns,
            default=default_columns,
            help="Sélectionnez les colonnes à inclure dans l'export"
        )
        
        # Filters
        st.markdown("##### 🔍 Filtres")
        
        filters = {}
        
        with st.expander("Filtres Avancés", expanded=False):
            # Date range filter
            if any('date' in col.lower() or 'period' in col.lower() for col in data.columns):
                col1, col2 = st.columns(2)
                with col1:
                    start_date = st.date_input(
                        "Date de début",
                        value=date.today() - timedelta(days=30)
                    )
                with col2:
                    end_date = st.date_input(
                        "Date de fin",
                        value=date.today()
                    )
                filters['date_range'] = (start_date, end_date)
            
            # Categorical filters
            categorical_columns = data.select_dtypes(include=['object']).columns
            for col in categorical_columns[:3]:  # Limit to first 3 categorical columns
                unique_values = data[col].dropna().unique()
                if len(unique_values) <= 20:  # Only show filter if manageable number of options
                    selected_values = st.multiselect(
                        f"Filtrer par {col}",
                        unique_values,
                        default=unique_values.tolist()
                    )
                    if len(selected_values) < len(unique_values):
                        filters[col] = selected_values
        
        # Export options
        st.markdown("##### ⚙️ Options d'Export")
        
        export_options = {}
        
        col1, col2, col3 = st.columns(3)
        
        with col1:
            export_options['include_index'] = st.checkbox(
                "Inclure les index",
                value=False,
                help="Inclure les numéros de ligne dans l'export"
            )
        
        with col2:
            export_options['include_header'] = st.checkbox(
                "Inclure les en-têtes",
                value=True,
                help="Inclure les noms de colonnes"
            )
        
        with col3:
            if format_name in ['Excel', 'PDF']:
                export_options['include_formatting'] = st.checkbox(
                    "Mise en forme avancée",
                    value=True,
                    help="Inclure couleurs, bordures et styles"
                )
        
        # Preview
        if selected_columns:
            self._show_export_preview(data, selected_columns, filters, format_name)
        
        # Export button
        if st.button("🚀 Exporter", type="primary", disabled=not selected_columns):
            export_config = {
                'columns': selected_columns,
                'filters': filters,
                'options': export_options
            }
            
            filename = f"export_personnalise_{data_type}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
            self._process_export(data, format_name, export_config, filename)
    
    def _create_scheduled_export(self, data: pd.DataFrame, data_type: str):
        """Create scheduled export interface"""
        
        st.markdown("#### 📅 Export Programmé")
        st.markdown("Programmez des exports automatiques à intervalles réguliers")
        
        # Current scheduled exports
        if st.session_state.scheduled_exports:
            st.markdown("##### 📋 Exports Programmés Actifs")
            
            for i, scheduled in enumerate(st.session_state.scheduled_exports):
                with st.expander(f"🕒 {scheduled['name']} - {scheduled['frequency']}"):
                    col1, col2, col3 = st.columns(3)
                    
                    with col1:
                        st.write(f"**Format:** {scheduled['format']}")
                        st.write(f"**Fréquence:** {scheduled['frequency']}")
                    
                    with col2:
                        st.write(f"**Prochaine exécution:** {scheduled['next_run']}")
                        st.write(f"**Destinataires:** {len(scheduled.get('recipients', []))}")
                    
                    with col3:
                        if st.button("❌ Supprimer", key=f"delete_scheduled_{i}"):
                            st.session_state.scheduled_exports.pop(i)
                            st.rerun()
                        
                        if st.button("▶️ Exécuter Maintenant", key=f"run_scheduled_{i}"):
                            st.info("Export en cours d'exécution...")
        else:
            st.info("Aucun export programmé actuellement")
        
        # Schedule new export
        st.markdown("##### ➕ Programmer un Nouvel Export")
        
        with st.form("schedule_export_form"):
            col1, col2 = st.columns(2)
            
            with col1:
                schedule_name = st.text_input(
                    "Nom de l'export",
                    placeholder="Ex: Rapport mensuel factures"
                )
                
                frequency = st.selectbox(
                    "Fréquence",
                    ["Quotidien", "Hebdomadaire", "Mensuel", "Trimestriel"],
                    help="À quelle fréquence exécuter l'export"
                )
                
                export_format = st.selectbox(
                    "Format",
                    [name for name, info in self.supported_formats.items() if info['available']]
                )
            
            with col2:
                start_time = st.time_input(
                    "Heure d'exécution",
                    value=datetime.strptime("09:00", "%H:%M").time(),
                    help="Heure d'exécution quotidienne"
                )
                
                recipients = st.text_area(
                    "Destinataires (emails)",
                    placeholder="email1@example.com, email2@example.com",
                    help="Emails séparés par des virgules"
                )
                
                include_charts = st.checkbox(
                    "Inclure graphiques",
                    value=True,
                    help="Inclure des graphiques dans l'export"
                )
            
            if st.form_submit_button("📅 Programmer l'Export", type="primary"):
                if schedule_name and recipients:
                    scheduled_export = {
                        'name': schedule_name,
                        'frequency': frequency,
                        'format': export_format,
                        'start_time': start_time.strftime("%H:%M"),
                        'recipients': [email.strip() for email in recipients.split(',')],
                        'include_charts': include_charts,
                        'created': datetime.now(),
                        'next_run': self._calculate_next_run(frequency, start_time),
                        'data_type': data_type
                    }
                    
                    st.session_state.scheduled_exports.append(scheduled_export)
                    st.success(f"✅ Export '{schedule_name}' programmé avec succès!")
                    st.rerun()
                else:
                    st.error("Veuillez remplir le nom et au moins un destinataire")
    
    def _show_export_history(self):
        """Show export history"""
        
        st.markdown("#### 📊 Historique des Exports")
        
        if st.session_state.export_history:
            # Summary metrics
            col1, col2, col3, col4 = st.columns(4)
            
            total_exports = len(st.session_state.export_history)
            successful_exports = len([e for e in st.session_state.export_history if e['status'] == 'success'])
            last_export = max(st.session_state.export_history, key=lambda x: x['timestamp'])
            popular_format = max(set([e['format'] for e in st.session_state.export_history]), 
                               key=[e['format'] for e in st.session_state.export_history].count)
            
            with col1:
                st.metric("Total Exports", total_exports)
            with col2:
                st.metric("Taux de Succès", f"{(successful_exports/total_exports)*100:.1f}%")
            with col3:
                st.metric("Dernier Export", last_export['timestamp'].strftime("%d/%m %H:%M"))
            with col4:
                st.metric("Format Populaire", popular_format)
            
            # Export history table
            st.markdown("##### 📋 Historique Détaillé")
            
            history_data = []
            for export in sorted(st.session_state.export_history, key=lambda x: x['timestamp'], reverse=True):
                status_emoji = "✅" if export['status'] == 'success' else "❌"
                history_data.append({
                    'Statut': f"{status_emoji} {export['status'].title()}",
                    'Nom': export['filename'],
                    'Format': f"{self.supported_formats[export['format']]['icon']} {export['format']}",
                    'Taille': export.get('size', 'N/A'),
                    'Date': export['timestamp'].strftime("%d/%m/%Y %H:%M"),
                    'Durée': f"{export.get('duration', 0):.1f}s"
                })
            
            df_history = pd.DataFrame(history_data)
            
            # Search and filter
            col1, col2 = st.columns([2, 1])
            with col1:
                search_term = st.text_input("🔍 Rechercher dans l'historique")
            with col2:
                status_filter = st.selectbox("Filtrer par statut", ["Tous", "Success", "Error"])
            
            # Apply filters
            if search_term:
                mask = df_history.astype(str).apply(lambda x: x.str.contains(search_term, case=False, na=False)).any(axis=1)
                df_history = df_history[mask]
            
            if status_filter != "Tous":
                df_history = df_history[df_history['Statut'].str.contains(status_filter, case=False)]
            
            st.dataframe(df_history, use_container_width=True)
            
            # Bulk actions
            if len(df_history) > 0:
                col1, col2 = st.columns(2)
                with col1:
                    if st.button("🗑️ Vider l'Historique"):
                        st.session_state.export_history.clear()
                        st.success("Historique vidé")
                        st.rerun()
                
                with col2:
                    if st.button("📥 Exporter l'Historique"):
                        history_csv = df_history.to_csv(index=False)
                        st.download_button(
                            "📊 Télécharger CSV",
                            history_csv,
                            file_name=f"historique_exports_{datetime.now().strftime('%Y%m%d')}.csv",
                            mime="text/csv"
                        )
        else:
            st.info("Aucun export dans l'historique")
            st.markdown("*Effectuez votre premier export pour voir l'historique ici*")
    
    def _show_export_preview(self, data: pd.DataFrame, columns: List[str], 
                           filters: Dict[str, Any], format_name: str):
        """Show preview of export data"""
        
        st.markdown("##### 👀 Aperçu de l'Export")
        
        # Apply filters
        filtered_data = self._apply_filters(data, filters)
        preview_data = filtered_data[columns]
        
        # Preview statistics
        col1, col2, col3, col4 = st.columns(4)
        
        with col1:
            st.metric("Lignes", len(preview_data))
        with col2:
            st.metric("Colonnes", len(columns))
        with col3:
            estimated_size = self._estimate_file_size(preview_data, format_name)
            st.metric("Taille Estimée", estimated_size)
        with col4:
            st.metric("Format", f"{self.supported_formats[format_name]['icon']} {format_name}")
        
        # Data preview
        st.markdown("**Aperçu des données (10 premières lignes):**")
        st.dataframe(preview_data.head(10), use_container_width=True)
        
        if len(preview_data) > 10:
            st.info(f"... et {len(preview_data) - 10} lignes supplémentaires")
    
    def _apply_filters(self, data: pd.DataFrame, filters: Dict[str, Any]) -> pd.DataFrame:
        """Apply filters to data"""
        filtered_data = data.copy()
        
        for filter_name, filter_value in filters.items():
            if filter_name == 'date_range' and filter_value:
                start_date, end_date = filter_value
                # Try to find date columns
                date_columns = [col for col in data.columns if 'date' in col.lower() or 'period' in col.lower()]
                if date_columns:
                    date_col = date_columns[0]
                    try:
                        filtered_data[date_col] = pd.to_datetime(filtered_data[date_col])
                        mask = (filtered_data[date_col].dt.date >= start_date) & (filtered_data[date_col].dt.date <= end_date)
                        filtered_data = filtered_data[mask]
                    except:
                        pass  # Skip if date conversion fails
            
            elif filter_name in data.columns and filter_value:
                filtered_data = filtered_data[filtered_data[filter_name].isin(filter_value)]
        
        return filtered_data
    
    def _estimate_file_size(self, data: pd.DataFrame, format_name: str) -> str:
        """Estimate file size based on data and format"""
        
        # Basic size estimation
        rows, cols = data.shape
        
        if format_name == "CSV":
            # Assume average 10 characters per cell
            size_bytes = rows * cols * 10
        elif format_name == "JSON":
            # JSON is usually larger due to structure
            size_bytes = rows * cols * 15
        elif format_name == "Excel":
            # Excel files are larger due to formatting
            size_bytes = rows * cols * 12
        elif format_name == "PDF":
            # PDF size depends on page count
            pages = max(1, rows // 30)  # Assume 30 rows per page
            size_bytes = pages * 50000  # 50KB per page
        else:
            size_bytes = rows * cols * 10
        
        # Convert to human readable
        if size_bytes < 1024:
            return f"{size_bytes} B"
        elif size_bytes < 1024 * 1024:
            return f"{size_bytes / 1024:.1f} KB"
        else:
            return f"{size_bytes / (1024 * 1024):.1f} MB"
    
    def _process_export(self, data: pd.DataFrame, format_name: str, 
                       config: Dict[str, Any], filename: str):
        """Process export request"""
        
        start_time = datetime.now()
        
        try:
            # Apply configuration
            export_data = data.copy()
            
            if 'columns' in config and config['columns']:
                export_data = export_data[config['columns']]
            
            if 'filters' in config:
                export_data = self._apply_filters(export_data, config['filters'])
            
            # Generate export based on format
            if format_name == "CSV":
                result = self._export_csv(export_data, config.get('options', {}))
            elif format_name == "JSON":
                result = self._export_json(export_data, config.get('options', {}))
            elif format_name == "Excel" and EXCEL_AVAILABLE:
                result = self._export_excel(export_data, config.get('options', {}))
            elif format_name == "PDF" and PDF_AVAILABLE:
                result = self._export_pdf(export_data, config.get('options', {}))
            elif format_name == "HTML":
                result = self._export_html(export_data, config.get('options', {}))
            else:
                st.error(f"Format {format_name} non supporté ou dépendances manquantes")
                return
            
            # Calculate duration
            duration = (datetime.now() - start_time).total_seconds()
            
            # Create download button
            file_extension = self.supported_formats[format_name]['extension']
            mime_type = self.supported_formats[format_name]['mime_type']
            
            st.download_button(
                f"📥 Télécharger {format_name}",
                result,
                file_name=f"{filename}{file_extension}",
                mime=mime_type,
                help=f"Télécharger le fichier au format {format_name}"
            )
            
            # Add to history
            export_entry = {
                'filename': filename,
                'format': format_name,
                'status': 'success',
                'timestamp': start_time,
                'duration': duration,
                'size': self._estimate_file_size(export_data, format_name),
                'rows': len(export_data),
                'columns': len(export_data.columns)
            }
            
            st.session_state.export_history.append(export_entry)
            
            st.success(f"✅ Export {format_name} généré avec succès! ({len(export_data)} lignes, {duration:.1f}s)")
            
        except Exception as e:
            # Add error to history
            export_entry = {
                'filename': filename,
                'format': format_name,
                'status': 'error',
                'timestamp': start_time,
                'error': str(e),
                'duration': (datetime.now() - start_time).total_seconds()
            }
            
            st.session_state.export_history.append(export_entry)
            st.error(f"❌ Erreur lors de l'export: {e}")
    
    def _export_csv(self, data: pd.DataFrame, options: Dict[str, Any]) -> bytes:
        """Export data as CSV"""
        output = io.StringIO()
        data.to_csv(
            output,
            index=options.get('include_index', False),
            header=options.get('include_header', True),
            encoding='utf-8'
        )
        return output.getvalue().encode('utf-8')
    
    def _export_json(self, data: pd.DataFrame, options: Dict[str, Any]) -> bytes:
        """Export data as JSON"""
        # Convert data to JSON
        if options.get('include_index', False):
            json_data = data.to_dict('index')
        else:
            json_data = data.to_dict('records')
        
        # Add metadata
        export_data = {
            'metadata': {
                'exported_at': datetime.now().isoformat(),
                'total_rows': len(data),
                'total_columns': len(data.columns),
                'columns': list(data.columns)
            },
            'data': json_data
        }
        
        return json.dumps(export_data, indent=2, default=str, ensure_ascii=False).encode('utf-8')
    
    def _export_excel(self, data: pd.DataFrame, options: Dict[str, Any]) -> bytes:
        """Export data as Excel with formatting"""
        output = io.BytesIO()
        
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            data.to_excel(
                writer,
                sheet_name='Données',
                index=options.get('include_index', False),
                header=options.get('include_header', True)
            )
            
            # Apply formatting if requested
            if options.get('include_formatting', True):
                workbook = writer.book
                worksheet = writer.sheets['Données']
                
                # Header formatting
                if options.get('include_header', True):
                    header_font = Font(bold=True, color="FFFFFF")
                    header_fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
                    
                    for cell in worksheet[1]:
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = Alignment(horizontal="center")
                
                # Auto-adjust column widths
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    
                    adjusted_width = min(max_length + 2, 50)
                    worksheet.column_dimensions[column_letter].width = adjusted_width
                
                # Add borders
                thin_border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )
                
                for row in worksheet.iter_rows():
                    for cell in row:
                        cell.border = thin_border
        
        return output.getvalue()
    
    def _export_pdf(self, data: pd.DataFrame, options: Dict[str, Any]) -> bytes:
        """Export data as PDF"""
        output = io.BytesIO()
        
        # Create PDF document
        doc = SimpleDocTemplate(output, pagesize=A4)
        elements = []
        
        # Styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=16,
            spaceAfter=30,
            textColor=colors.HexColor('#2E7D32')
        )
        
        # Title
        title = Paragraph("Export de Données - Facturation PMO", title_style)
        elements.append(title)
        elements.append(Spacer(1, 12))
        
        # Metadata
        metadata_text = f"""
        <b>Date d'export:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}<br/>
        <b>Nombre de lignes:</b> {len(data)}<br/>
        <b>Nombre de colonnes:</b> {len(data.columns)}<br/>
        """
        metadata = Paragraph(metadata_text, styles['Normal'])
        elements.append(metadata)
        elements.append(Spacer(1, 20))
        
        # Data table
        if not data.empty:
            # Limit data for PDF (max 50 rows to avoid huge files)
            display_data = data.head(50)
            
            # Prepare table data
            table_data = []
            
            if options.get('include_header', True):
                table_data.append(list(display_data.columns))
            
            for _, row in display_data.iterrows():
                table_data.append([str(value)[:30] + '...' if len(str(value)) > 30 else str(value) 
                                 for value in row])
            
            # Create table
            table = Table(table_data)
            
            # Table style
            table_style = TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E7D32')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ])
            
            table.setStyle(table_style)
            elements.append(table)
            
            if len(data) > 50:
                elements.append(Spacer(1, 12))
                note = Paragraph(f"<i>Note: Seules les 50 premières lignes sont affichées. Total: {len(data)} lignes.</i>", 
                               styles['Normal'])
                elements.append(note)
        
        # Build PDF
        doc.build(elements)
        return output.getvalue()
    
    def _export_html(self, data: pd.DataFrame, options: Dict[str, Any]) -> bytes:
        """Export data as HTML"""
        # Generate HTML
        html_content = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <meta charset="utf-8">
            <title>Export de Données - Facturation PMO</title>
            <style>
                body {{
                    font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
                    margin: 40px;
                    background-color: #f5f5f5;
                }}
                .header {{
                    background: linear-gradient(135deg, #2E7D32, #4CAF50);
                    color: white;
                    padding: 20px;
                    border-radius: 8px;
                    margin-bottom: 20px;
                }}
                .metadata {{
                    background: white;
                    padding: 15px;
                    border-radius: 8px;
                    margin-bottom: 20px;
                    box-shadow: 0 2px 4px rgba(0,0,0,0.1);
                }}
                table {{
                    width: 100%;
                    border-collapse: collapse;
                    background: white;
                    border-radius: 8px;
                    overflow: hidden;
                    box-shadow: 0 2px 4px rgba(0,0,0,0.1);
                }}
                th {{
                    background: #2E7D32;
                    color: white;
                    padding: 12px;
                    text-align: left;
                    font-weight: 600;
                }}
                td {{
                    padding: 10px 12px;
                    border-bottom: 1px solid #e0e0e0;
                }}
                tr:nth-child(even) {{
                    background-color: #f9f9f9;
                }}
                tr:hover {{
                    background-color: #e8f5e9;
                }}
                .footer {{
                    margin-top: 20px;
                    text-align: center;
                    color: #666;
                    font-size: 0.9em;
                }}
            </style>
        </head>
        <body>
            <div class="header">
                <h1>Export de Données - Facturation PMO</h1>
                <p>Système de facturation pour l'autoconsommation collective photovoltaïque</p>
            </div>
            
            <div class="metadata">
                <h3>Informations de l'Export</h3>
                <p><strong>Date d'export:</strong> {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}</p>
                <p><strong>Nombre de lignes:</strong> {len(data)}</p>
                <p><strong>Nombre de colonnes:</strong> {len(data.columns)}</p>
                <p><strong>Colonnes:</strong> {', '.join(data.columns)}</p>
            </div>
            
            {data.to_html(index=options.get('include_index', False), 
                         table_id='data-table',
                         classes='table table-striped',
                         escape=False)}
            
            <div class="footer">
                <p>Généré par OptimPV - Système de Facturation PMO</p>
                <p>© 2025 OptimPV. Tous droits réservés.</p>
            </div>
        </body>
        </html>
        """
        
        return html_content.encode('utf-8')
    
    def _calculate_next_run(self, frequency: str, start_time) -> datetime:
        """Calculate next run time for scheduled export"""
        now = datetime.now()
        today = now.date()
        
        # Combine today's date with the scheduled time
        next_run = datetime.combine(today, start_time)
        
        # If the time has already passed today, schedule for next occurrence
        if next_run <= now:
            if frequency == "Quotidien":
                next_run += timedelta(days=1)
            elif frequency == "Hebdomadaire":
                next_run += timedelta(weeks=1)
            elif frequency == "Mensuel":
                # Add one month (approximately)
                next_run += timedelta(days=30)
            elif frequency == "Trimestriel":
                next_run += timedelta(days=90)
        
        return next_run
    
    def _save_export_template(self):
        """Save current export configuration as template"""
        with st.sidebar:
            with st.form("save_template_form"):
                st.markdown("#### 💾 Sauvegarder le Modèle")
                
                template_name = st.text_input("Nom du modèle")
                template_description = st.text_area("Description")
                
                if st.form_submit_button("Sauvegarder"):
                    if template_name:
                        # This would capture current form state in a real implementation
                        template_config = {
                            'description': template_description,
                            'columns': ['all'],  # Placeholder
                            'filters': {},
                            'format': 'CSV'
                        }
                        
                        st.session_state.export_templates[template_name] = template_config
                        st.success(f"✅ Modèle '{template_name}' sauvegardé!")
                    else:
                        st.error("Veuillez entrer un nom pour le modèle")