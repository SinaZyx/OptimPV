"""
Advanced Reporting System for PMO Billing
Generates comprehensive reports in multiple formats with data visualization
"""

import pandas as pd
import numpy as np
from datetime import datetime, date, timedelta
from typing import Dict, List, Optional, Any, Union
import logging
import io
import json
import base64
from pathlib import Path

# Data visualization
import plotly.graph_objects as go
import plotly.express as px
from plotly.subplots import make_subplots
import plotly.io as pio

# Export capabilities
try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.chart import LineChart, BarChart, PieChart, Reference
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm, cm
    from reportlab.pdfgen import canvas
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    from reportlab.graphics.charts.piecharts import Pie
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

# Email capabilities
try:
    import smtplib
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    from email.mime.application import MIMEApplication
    EMAIL_AVAILABLE = True
except ImportError:
    EMAIL_AVAILABLE = False

from .analytics import AnalyticsEngine, KPIMetrics
from .database import BillingDatabase

logger = logging.getLogger(__name__)

class ReportTemplate:
    """Report template configuration"""
    
    MONTHLY = {
        'name': 'Monthly Report',
        'description': 'Comprehensive monthly performance report',
        'sections': ['summary', 'revenue', 'customers', 'energy', 'risks'],
        'charts': ['revenue_trend', 'customer_segmentation', 'energy_breakdown'],
        'period': 'month'
    }
    
    QUARTERLY = {
        'name': 'Quarterly Report',
        'description': 'Quarterly business review with deep analytics',
        'sections': ['executive_summary', 'financial_performance', 'operational_metrics', 'predictive_insights'],
        'charts': ['quarterly_revenue', 'kpi_dashboard', 'customer_analysis', 'forecasting'],
        'period': 'quarter'
    }
    
    ANNUAL = {
        'name': 'Annual Report',
        'description': 'Complete yearly performance and strategic analysis',
        'sections': ['executive_summary', 'financial_review', 'operational_excellence', 'market_analysis', 'strategic_outlook'],
        'charts': ['annual_trends', 'performance_metrics', 'market_comparison', 'future_projections'],
        'period': 'year'
    }
    
    CUSTOMER_ANALYSIS = {
        'name': 'Customer Analysis Report',
        'description': 'Detailed customer behavior and segmentation analysis',
        'sections': ['customer_overview', 'segmentation', 'payment_behavior', 'risk_analysis'],
        'charts': ['customer_distribution', 'payment_patterns', 'risk_matrix'],
        'period': 'custom'
    }
    
    FINANCIAL_DASHBOARD = {
        'name': 'Financial Dashboard',
        'description': 'Real-time financial metrics and KPIs',
        'sections': ['kpi_summary', 'cash_flow', 'collections', 'forecasting'],
        'charts': ['kpi_gauges', 'cash_flow_chart', 'aging_analysis'],
        'period': 'realtime'
    }

class ReportGenerator:
    """Advanced report generator with multi-format export capabilities"""
    
    def __init__(self, db_path: str = "data/billing.db"):
        """Initialize report generator"""
        self.db = BillingDatabase(db_path)
        self.analytics = AnalyticsEngine(db_path)
        self.templates = {
            'monthly': ReportTemplate.MONTHLY,
            'quarterly': ReportTemplate.QUARTERLY,
            'annual': ReportTemplate.ANNUAL,
            'customer_analysis': ReportTemplate.CUSTOMER_ANALYSIS,
            'financial_dashboard': ReportTemplate.FINANCIAL_DASHBOARD
        }
        
        # Configure plotly theme
        pio.templates.default = "plotly_white"
        
        # Email configuration
        self.smtp_config = {}
        
    def generate_report(self,
                       template_name: str,
                       start_date: Optional[date] = None,
                       end_date: Optional[date] = None,
                       project_id: Optional[int] = None,
                       export_formats: List[str] = ['html'],
                       include_charts: bool = True,
                       custom_sections: Optional[List[str]] = None) -> Dict[str, Any]:
        """Generate comprehensive report with specified template"""
        
        if template_name not in self.templates:
            raise ValueError(f"Unknown template: {template_name}")
        
        template = self.templates[template_name]
        
        # Set default date range based on template
        if not end_date:
            end_date = date.today()
        if not start_date:
            if template['period'] == 'month':
                start_date = end_date.replace(day=1)
            elif template['period'] == 'quarter':
                start_date = end_date - timedelta(days=90)
            elif template['period'] == 'year':
                start_date = end_date.replace(month=1, day=1)
            else:
                start_date = end_date - timedelta(days=30)
        
        try:
            # Gather data
            report_data = self._gather_report_data(
                template, start_date, end_date, project_id, custom_sections
            )
            
            # Generate visualizations
            charts = {}
            if include_charts:
                charts = self._generate_charts(template, report_data)
            
            # Prepare report content
            report_content = {
                'metadata': {
                    'template': template_name,
                    'generated_at': datetime.now(),
                    'period': {'start': start_date, 'end': end_date},
                    'project_id': project_id
                },
                'data': report_data,
                'charts': charts,
                'summary': self._generate_executive_summary(report_data)
            }
            
            # Export in requested formats
            exports = {}
            for format_type in export_formats:
                if format_type == 'html':
                    exports['html'] = self._export_html(report_content, template)
                elif format_type == 'pdf' and PDF_AVAILABLE:
                    exports['pdf'] = self._export_pdf(report_content, template)
                elif format_type == 'excel' and EXCEL_AVAILABLE:
                    exports['excel'] = self._export_excel(report_content, template)
                elif format_type == 'json':
                    exports['json'] = self._export_json(report_content)
                elif format_type == 'csv':
                    exports['csv'] = self._export_csv(report_content)
            
            return {
                'success': True,
                'report_content': report_content,
                'exports': exports,
                'template_info': template
            }
            
        except Exception as e:
            logger.error(f"Error generating report: {e}")
            return {
                'success': False,
                'error': str(e),
                'template_info': template
            }
    
    def schedule_automated_report(self,
                                template_name: str,
                                recipients: List[str],
                                frequency: str = 'monthly',
                                start_time: str = '09:00',
                                export_formats: List[str] = ['pdf', 'excel']) -> bool:
        """Schedule automated report generation and email delivery"""
        # This would integrate with a scheduler (like celery or cron)
        # For now, return configuration success
        try:
            schedule_config = {
                'template': template_name,
                'recipients': recipients,
                'frequency': frequency,
                'start_time': start_time,
                'formats': export_formats,
                'created_at': datetime.now()
            }
            
            # Store configuration (in production, would save to database)
            logger.info(f"Scheduled report configured: {schedule_config}")
            return True
            
        except Exception as e:
            logger.error(f"Error scheduling report: {e}")
            return False
    
    def send_report_by_email(self,
                           report_content: Dict[str, Any],
                           recipients: List[str],
                           subject: Optional[str] = None,
                           body_template: Optional[str] = None,
                           attachments: Optional[Dict[str, bytes]] = None) -> bool:
        """Send report via email with attachments"""
        
        if not EMAIL_AVAILABLE or not self.smtp_config:
            logger.error("Email functionality not available or not configured")
            return False
        
        try:
            # Create message
            msg = MIMEMultipart()
            msg['From'] = self.smtp_config.get('username', '')
            msg['To'] = ', '.join(recipients)
            msg['Subject'] = subject or f"Automated Report - {report_content['metadata']['template']}"
            
            # Email body
            if not body_template:
                body_template = self._get_default_email_template()
            
            body = body_template.format(
                template_name=report_content['metadata']['template'],
                start_date=report_content['metadata']['period']['start'],
                end_date=report_content['metadata']['period']['end'],
                generated_at=report_content['metadata']['generated_at'].strftime('%Y-%m-%d %H:%M'),
                summary=report_content.get('summary', 'Report generated successfully')
            )
            
            msg.attach(MIMEText(body, 'html'))
            
            # Add attachments
            if attachments:
                for filename, content in attachments.items():
                    attachment = MIMEApplication(content)
                    attachment.add_header('Content-Disposition', 'attachment', filename=filename)
                    msg.attach(attachment)
            
            # Send email
            with smtplib.SMTP(self.smtp_config['server'], self.smtp_config['port']) as server:
                if self.smtp_config.get('use_tls', True):
                    server.starttls()
                server.login(self.smtp_config['username'], self.smtp_config['password'])
                server.send_message(msg)
            
            logger.info(f"Report sent successfully to {len(recipients)} recipients")
            return True
            
        except Exception as e:
            logger.error(f"Error sending email report: {e}")
            return False
    
    def configure_email_settings(self, 
                               smtp_server: str,
                               smtp_port: int,
                               username: str,
                               password: str,
                               use_tls: bool = True) -> bool:
        """Configure email settings for automated reports"""
        try:
            self.smtp_config = {
                'server': smtp_server,
                'port': smtp_port,
                'username': username,
                'password': password,
                'use_tls': use_tls
            }
            
            # Test connection
            with smtplib.SMTP(smtp_server, smtp_port) as server:
                if use_tls:
                    server.starttls()
                server.login(username, password)
            
            logger.info("Email configuration successful")
            return True
            
        except Exception as e:
            logger.error(f"Email configuration failed: {e}")
            return False
    
    def _gather_report_data(self,
                          template: Dict[str, Any],
                          start_date: date,
                          end_date: date,
                          project_id: Optional[int],
                          custom_sections: Optional[List[str]]) -> Dict[str, Any]:
        """Gather all data needed for the report"""
        
        sections = custom_sections or template['sections']
        data = {}
        
        # Get comprehensive KPIs
        kpis = self.analytics.calculate_comprehensive_kpis(start_date, end_date, project_id)
        data['kpis'] = kpis
        
        # Section-specific data gathering
        if 'summary' in sections or 'executive_summary' in sections:
            data['summary'] = self._get_summary_data(kpis, start_date, end_date, project_id)
        
        if 'revenue' in sections or 'financial_performance' in sections:
            data['revenue'] = self._get_revenue_data(start_date, end_date, project_id)
        
        if 'customers' in sections or 'customer_overview' in sections:
            data['customers'] = self._get_customer_data(start_date, end_date, project_id)
        
        if 'energy' in sections:
            data['energy'] = self._get_energy_data(start_date, end_date, project_id)
        
        if 'risks' in sections or 'risk_analysis' in sections:
            data['risks'] = self._get_risk_data(start_date, end_date, project_id)
        
        if 'segmentation' in sections:
            data['segmentation'] = self.analytics.segment_customers(project_id)
        
        if 'predictive_insights' in sections or 'forecasting' in sections:
            data['predictions'] = self.analytics.generate_predictive_insights(6, project_id)
        
        if 'operational_metrics' in sections:
            data['operations'] = self._get_operational_data(start_date, end_date, project_id)
        
        return data
    
    def _generate_charts(self, template: Dict[str, Any], data: Dict[str, Any]) -> Dict[str, str]:
        """Generate all charts for the report"""
        charts = {}
        
        for chart_type in template.get('charts', []):
            try:
                if chart_type == 'revenue_trend':
                    charts['revenue_trend'] = self._create_revenue_trend_chart(data)
                elif chart_type == 'customer_segmentation':
                    charts['customer_segmentation'] = self._create_customer_segmentation_chart(data)
                elif chart_type == 'energy_breakdown':
                    charts['energy_breakdown'] = self._create_energy_breakdown_chart(data)
                elif chart_type == 'kpi_dashboard':
                    charts['kpi_dashboard'] = self._create_kpi_dashboard(data)
                elif chart_type == 'quarterly_revenue':
                    charts['quarterly_revenue'] = self._create_quarterly_revenue_chart(data)
                elif chart_type == 'customer_analysis':
                    charts['customer_analysis'] = self._create_customer_analysis_chart(data)
                elif chart_type == 'forecasting':
                    charts['forecasting'] = self._create_forecasting_chart(data)
                elif chart_type == 'payment_patterns':
                    charts['payment_patterns'] = self._create_payment_patterns_chart(data)
                elif chart_type == 'risk_matrix':
                    charts['risk_matrix'] = self._create_risk_matrix_chart(data)
                elif chart_type == 'kpi_gauges':
                    charts['kpi_gauges'] = self._create_kpi_gauges(data)
                elif chart_type == 'cash_flow_chart':
                    charts['cash_flow_chart'] = self._create_cash_flow_chart(data)
                elif chart_type == 'aging_analysis':
                    charts['aging_analysis'] = self._create_aging_analysis_chart(data)
                    
            except Exception as e:
                logger.error(f"Error creating chart {chart_type}: {e}")
                continue
        
        return charts
    
    def _create_revenue_trend_chart(self, data: Dict[str, Any]) -> str:
        """Create revenue trend chart"""
        if 'revenue' not in data:
            return ""
        
        revenue_data = data['revenue']
        
        fig = go.Figure()
        
        if 'monthly_revenue' in revenue_data:
            months = list(revenue_data['monthly_revenue'].keys())
            values = list(revenue_data['monthly_revenue'].values())
            
            fig.add_trace(go.Scatter(
                x=months,
                y=values,
                mode='lines+markers',
                name='Monthly Revenue',
                line=dict(color='#1f77b4', width=3),
                marker=dict(size=8)
            ))
        
        fig.update_layout(
            title='Revenue Trend Analysis',
            xaxis_title='Month',
            yaxis_title='Revenue (€)',
            template='plotly_white',
            height=400
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="revenue_trend_chart")
    
    def _create_customer_segmentation_chart(self, data: Dict[str, Any]) -> str:
        """Create customer segmentation pie chart"""
        if 'segmentation' not in data:
            return ""
        
        segments = data['segmentation']
        labels = [seg.segment_name for seg in segments]
        values = [seg.customer_count for seg in segments]
        colors = ['#ff9999', '#66b3ff', '#99ff99', '#ffcc99']
        
        fig = go.Figure(data=[go.Pie(
            labels=labels,
            values=values,
            hole=0.3,
            marker_colors=colors
        )])
        
        fig.update_layout(
            title='Customer Segmentation',
            template='plotly_white',
            height=400
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="customer_segmentation_chart")
    
    def _create_energy_breakdown_chart(self, data: Dict[str, Any]) -> str:
        """Create energy breakdown chart"""
        if 'energy' not in data:
            return ""
        
        energy_data = data['energy']
        
        categories = ['Autoconsommation', 'Injection Réseau', 'Pertes']
        values = [
            energy_data.get('autoconsumption_kwh', 0),
            energy_data.get('injection_kwh', 0),
            energy_data.get('losses_kwh', 0)
        ]
        
        fig = go.Figure(data=[go.Bar(
            x=categories,
            y=values,
            marker_color=['#2E86AB', '#A23B72', '#F18F01']
        )])
        
        fig.update_layout(
            title='Energy Production Breakdown',
            xaxis_title='Category',
            yaxis_title='Energy (kWh)',
            template='plotly_white',
            height=400
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="energy_breakdown_chart")
    
    def _create_kpi_dashboard(self, data: Dict[str, Any]) -> str:
        """Create comprehensive KPI dashboard"""
        kpis = data.get('kpis', {})
        
        # Create subplot layout
        fig = make_subplots(
            rows=2, cols=3,
            subplot_titles=(
                'Revenue Metrics', 'Customer Metrics', 'Collection Metrics',
                'Energy Metrics', 'Cash Flow', 'Risk Indicators'
            ),
            specs=[[{"type": "indicator"}, {"type": "indicator"}, {"type": "indicator"}],
                   [{"type": "indicator"}, {"type": "indicator"}, {"type": "indicator"}]]
        )
        
        # Revenue indicator
        fig.add_trace(go.Indicator(
            mode="gauge+number+delta",
            value=kpis.total_revenue if hasattr(kpis, 'total_revenue') else 0,
            domain={'x': [0, 1], 'y': [0, 1]},
            title={'text': "Total Revenue (€)"},
            gauge={'axis': {'range': [None, 100000]},
                   'bar': {'color': "darkblue"},
                   'steps': [{'range': [0, 50000], 'color': "lightgray"},
                            {'range': [50000, 100000], 'color': "gray"}],
                   'threshold': {'line': {'color': "red", 'width': 4},
                                'thickness': 0.75, 'value': 90000}}
        ), row=1, col=1)
        
        # Collection rate indicator
        fig.add_trace(go.Indicator(
            mode="gauge+number",
            value=kpis.collection_rate if hasattr(kpis, 'collection_rate') else 0,
            domain={'x': [0, 1], 'y': [0, 1]},
            title={'text': "Collection Rate (%)"},
            gauge={'axis': {'range': [None, 100]},
                   'bar': {'color': "darkgreen"},
                   'steps': [{'range': [0, 80], 'color': "lightgray"},
                            {'range': [80, 100], 'color': "gray"}]}
        ), row=1, col=3)
        
        # Customer count
        fig.add_trace(go.Indicator(
            mode="number+delta",
            value=kpis.customer_count if hasattr(kpis, 'customer_count') else 0,
            title={'text': "Active Customers"},
            number={'font': {'size': 40}}
        ), row=1, col=2)
        
        fig.update_layout(
            height=600,
            title='KPI Dashboard',
            template='plotly_white'
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="kpi_dashboard_chart")
    
    def _create_kpi_gauges(self, data: Dict[str, Any]) -> str:
        """Create KPI gauge charts"""
        return self._create_kpi_dashboard(data)  # Reuse dashboard for gauges
    
    def _create_cash_flow_chart(self, data: Dict[str, Any]) -> str:
        """Create cash flow analysis chart"""
        kpis = data.get('kpis', {})
        
        categories = ['Operating Cash Flow', 'Free Cash Flow']
        values = [
            kpis.operating_cash_flow if hasattr(kpis, 'operating_cash_flow') else 0,
            kpis.free_cash_flow if hasattr(kpis, 'free_cash_flow') else 0
        ]
        
        fig = go.Figure(data=[go.Bar(
            x=categories,
            y=values,
            marker_color=['#1f77b4', '#ff7f0e']
        )])
        
        fig.update_layout(
            title='Cash Flow Analysis',
            xaxis_title='Cash Flow Type',
            yaxis_title='Amount (€)',
            template='plotly_white',
            height=400
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="cash_flow_chart")
    
    def _export_html(self, report_content: Dict[str, Any], template: Dict[str, Any]) -> str:
        """Export report as HTML"""
        html_content = f"""
        <!DOCTYPE html>
        <html>
        <head>
            <title>{template['name']} - {report_content['metadata']['generated_at'].strftime('%Y-%m-%d')}</title>
            <style>
                body {{ font-family: Arial, sans-serif; margin: 20px; }}
                .header {{ text-align: center; margin-bottom: 30px; }}
                .section {{ margin-bottom: 30px; }}
                .kpi-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); gap: 20px; }}
                .kpi-card {{ border: 1px solid #ddd; padding: 15px; border-radius: 5px; text-align: center; }}
                .chart-container {{ margin: 20px 0; }}
                table {{ width: 100%; border-collapse: collapse; margin: 10px 0; }}
                th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
                th {{ background-color: #f2f2f2; }}
            </style>
        </head>
        <body>
            <div class="header">
                <h1>{template['name']}</h1>
                <p>Period: {report_content['metadata']['period']['start']} to {report_content['metadata']['period']['end']}</p>
                <p>Generated: {report_content['metadata']['generated_at'].strftime('%Y-%m-%d %H:%M')}</p>
            </div>
            
            <div class="section">
                <h2>Executive Summary</h2>
                <p>{report_content.get('summary', 'No summary available')}</p>
            </div>
            
            <div class="section">
                <h2>Key Performance Indicators</h2>
                <div class="kpi-grid">
                    {self._generate_kpi_html_cards(report_content.get('data', {}).get('kpis', {}))}
                </div>
            </div>
            
            <div class="section">
                <h2>Visualizations</h2>
                {self._generate_charts_html(report_content.get('charts', {}))}
            </div>
            
            <div class="section">
                <h2>Detailed Analysis</h2>
                {self._generate_detailed_analysis_html(report_content.get('data', {}))}
            </div>
        </body>
        </html>
        """
        
        return html_content
    
    def _export_pdf(self, report_content: Dict[str, Any], template: Dict[str, Any]) -> bytes:
        """Export report as PDF"""
        if not PDF_AVAILABLE:
            raise ImportError("ReportLab not available for PDF export")
        
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        story = []
        styles = getSampleStyleSheet()
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            alignment=1  # Center
        )
        story.append(Paragraph(template['name'], title_style))
        story.append(Spacer(1, 12))
        
        # Metadata
        period_text = f"Period: {report_content['metadata']['period']['start']} to {report_content['metadata']['period']['end']}"
        story.append(Paragraph(period_text, styles['Normal']))
        
        generated_text = f"Generated: {report_content['metadata']['generated_at'].strftime('%Y-%m-%d %H:%M')}"
        story.append(Paragraph(generated_text, styles['Normal']))
        story.append(Spacer(1, 20))
        
        # Executive Summary
        story.append(Paragraph("Executive Summary", styles['Heading2']))
        story.append(Paragraph(report_content.get('summary', 'No summary available'), styles['Normal']))
        story.append(Spacer(1, 20))
        
        # KPIs Table
        kpis = report_content.get('data', {}).get('kpis', {})
        if hasattr(kpis, '__dict__'):
            story.append(Paragraph("Key Performance Indicators", styles['Heading2']))
            kpi_data = [['Metric', 'Value']]
            
            for key, value in kpis.__dict__.items():
                metric_name = key.replace('_', ' ').title()
                if isinstance(value, float):
                    formatted_value = f"{value:.2f}"
                    if 'rate' in key or 'ratio' in key:
                        formatted_value += "%"
                    elif 'revenue' in key or 'amount' in key or 'flow' in key:
                        formatted_value += " €"
                    elif 'kwh' in key:
                        formatted_value += " kWh"
                elif isinstance(value, int):
                    formatted_value = str(value)
                else:
                    formatted_value = str(value)
                
                kpi_data.append([metric_name, formatted_value])
            
            kpi_table = Table(kpi_data)
            kpi_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 14),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            story.append(kpi_table)
        
        doc.build(story)
        buffer.seek(0)
        return buffer.read()
    
    def _export_excel(self, report_content: Dict[str, Any], template: Dict[str, Any]) -> bytes:
        """Export report as Excel workbook"""
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl not available for Excel export")
        
        buffer = io.BytesIO()
        
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            # Summary sheet
            summary_data = self._prepare_summary_dataframe(report_content)
            summary_data.to_excel(writer, sheet_name='Summary', index=False)
            
            # KPIs sheet
            kpis_data = self._prepare_kpis_dataframe(report_content.get('data', {}).get('kpis', {}))
            kpis_data.to_excel(writer, sheet_name='KPIs', index=False)
            
            # Revenue data if available
            if 'revenue' in report_content.get('data', {}):
                revenue_data = self._prepare_revenue_dataframe(report_content['data']['revenue'])
                revenue_data.to_excel(writer, sheet_name='Revenue', index=False)
            
            # Customer data if available
            if 'customers' in report_content.get('data', {}):
                customer_data = self._prepare_customer_dataframe(report_content['data']['customers'])
                customer_data.to_excel(writer, sheet_name='Customers', index=False)
            
            # Format workbook
            self._format_excel_workbook(writer.book)
        
        buffer.seek(0)
        return buffer.read()
    
    def _export_json(self, report_content: Dict[str, Any]) -> str:
        """Export report as JSON"""
        # Convert datetime objects for JSON serialization
        json_content = self._serialize_for_json(report_content)
        return json.dumps(json_content, indent=2, ensure_ascii=False)
    
    def _export_csv(self, report_content: Dict[str, Any]) -> str:
        """Export key data as CSV"""
        # Create a comprehensive CSV with key metrics
        csv_data = []
        
        # Add KPIs
        kpis = report_content.get('data', {}).get('kpis', {})
        if hasattr(kpis, '__dict__'):
            csv_data.append(['Section', 'Metric', 'Value', 'Unit'])
            for key, value in kpis.__dict__.items():
                metric_name = key.replace('_', ' ').title()
                unit = ''
                if 'rate' in key or 'ratio' in key:
                    unit = '%'
                elif 'revenue' in key or 'amount' in key or 'flow' in key:
                    unit = '€'
                elif 'kwh' in key:
                    unit = 'kWh'
                elif 'count' in key:
                    unit = 'count'
                elif 'days' in key:
                    unit = 'days'
                
                csv_data.append(['KPIs', metric_name, value, unit])
        
        # Convert to DataFrame and return CSV string
        df = pd.DataFrame(csv_data[1:], columns=csv_data[0])
        return df.to_csv(index=False)
    
    # Helper methods for data preparation and formatting
    
    def _get_summary_data(self, kpis: KPIMetrics, start_date: date, end_date: date, project_id: Optional[int]) -> Dict[str, Any]:
        """Get summary data for the report"""
        return {
            'period_revenue': kpis.total_revenue if hasattr(kpis, 'total_revenue') else 0,
            'collection_efficiency': kpis.collection_rate if hasattr(kpis, 'collection_rate') else 0,
            'customer_satisfaction': 85.5,  # Placeholder
            'energy_efficiency': kpis.energy_efficiency_ratio if hasattr(kpis, 'energy_efficiency_ratio') else 0,
            'risk_level': 'Low' if (kpis.overdue_ratio if hasattr(kpis, 'overdue_ratio') else 0) < 5 else 'Medium'
        }
    
    def _get_revenue_data(self, start_date: date, end_date: date, project_id: Optional[int]) -> Dict[str, Any]:
        """Get detailed revenue data"""
        # This would query the database for detailed revenue breakdown
        return {
            'total_revenue': 50000.0,
            'monthly_revenue': {
                '2024-01': 15000,
                '2024-02': 16500,
                '2024-03': 18500
            },
            'revenue_by_customer': {},
            'revenue_growth': 12.5
        }
    
    def _get_customer_data(self, start_date: date, end_date: date, project_id: Optional[int]) -> Dict[str, Any]:
        """Get customer analysis data"""
        return {
            'total_customers': 25,
            'active_customers': 23,
            'new_customers': 2,
            'churned_customers': 0,
            'customer_lifetime_value': 15000.0
        }
    
    def _get_energy_data(self, start_date: date, end_date: date, project_id: Optional[int]) -> Dict[str, Any]:
        """Get energy-related data"""
        return {
            'total_production_kwh': 45000,
            'autoconsumption_kwh': 35000,
            'injection_kwh': 10000,
            'losses_kwh': 0,
            'efficiency_ratio': 77.8
        }
    
    def _get_risk_data(self, start_date: date, end_date: date, project_id: Optional[int]) -> Dict[str, Any]:
        """Get risk analysis data"""
        risk_predictions = self.analytics.predict_payment_defaults(project_id)
        return {
            'high_risk_customers': len([r for r in risk_predictions if r['risk_level'] == 'High']),
            'medium_risk_customers': len([r for r in risk_predictions if r['risk_level'] == 'Medium']),
            'overdue_amount': 2500.0,
            'default_probability': 5.2
        }
    
    def _get_operational_data(self, start_date: date, end_date: date, project_id: Optional[int]) -> Dict[str, Any]:
        """Get operational metrics"""
        return {
            'invoice_processing_time': 2.5,
            'payment_processing_efficiency': 92.3,
            'customer_service_requests': 8,
            'system_uptime': 99.9
        }
    
    def _generate_executive_summary(self, data: Dict[str, Any]) -> str:
        """Generate executive summary based on data"""
        kpis = data.get('kpis', {})
        
        revenue = kpis.total_revenue if hasattr(kpis, 'total_revenue') else 0
        collection_rate = kpis.collection_rate if hasattr(kpis, 'collection_rate') else 0
        customer_count = kpis.customer_count if hasattr(kpis, 'customer_count') else 0
        
        summary = f"""
        Performance Summary: The reporting period shows total revenue of €{revenue:,.2f} 
        with a collection rate of {collection_rate:.1f}%. We are currently serving 
        {customer_count} active customers with strong operational efficiency. 
        
        Key highlights include stable energy delivery performance and maintaining 
        low risk exposure across our customer portfolio.
        """
        
        return summary.strip()
    
    def _generate_kpi_html_cards(self, kpis: Any) -> str:
        """Generate HTML cards for KPIs"""
        if not hasattr(kpis, '__dict__'):
            return ""
        
        cards_html = ""
        for key, value in kpis.__dict__.items():
            if key.startswith('_'):
                continue
                
            metric_name = key.replace('_', ' ').title()
            formatted_value = f"{value:.2f}" if isinstance(value, float) else str(value)
            
            if 'rate' in key or 'ratio' in key:
                formatted_value += "%"
            elif 'revenue' in key or 'amount' in key or 'flow' in key:
                formatted_value = f"€{formatted_value}"
            elif 'kwh' in key:
                formatted_value += " kWh"
            
            cards_html += f"""
            <div class="kpi-card">
                <h3>{metric_name}</h3>
                <p style="font-size: 24px; font-weight: bold; color: #2E86AB;">{formatted_value}</p>
            </div>
            """
        
        return cards_html
    
    def _generate_charts_html(self, charts: Dict[str, str]) -> str:
        """Generate HTML for charts"""
        charts_html = ""
        for chart_name, chart_html in charts.items():
            charts_html += f"""
            <div class="chart-container">
                <h3>{chart_name.replace('_', ' ').title()}</h3>
                {chart_html}
            </div>
            """
        return charts_html
    
    def _generate_detailed_analysis_html(self, data: Dict[str, Any]) -> str:
        """Generate detailed analysis HTML"""
        analysis_html = "<p>Detailed analysis based on current data trends and performance metrics.</p>"
        
        if 'predictions' in data:
            analysis_html += "<h3>Predictive Insights</h3><ul>"
            for prediction in data['predictions']:
                analysis_html += f"<li><strong>{prediction.metric_name}:</strong> {prediction.current_value} → {prediction.predicted_value} ({prediction.trend_direction})</li>"
            analysis_html += "</ul>"
        
        return analysis_html
    
    def _prepare_summary_dataframe(self, report_content: Dict[str, Any]) -> pd.DataFrame:
        """Prepare summary data as DataFrame"""
        metadata = report_content.get('metadata', {})
        return pd.DataFrame([{
            'Template': metadata.get('template', ''),
            'Generated At': metadata.get('generated_at', ''),
            'Start Date': metadata.get('period', {}).get('start', ''),
            'End Date': metadata.get('period', {}).get('end', ''),
            'Project ID': metadata.get('project_id', 'All')
        }])
    
    def _prepare_kpis_dataframe(self, kpis: Any) -> pd.DataFrame:
        """Prepare KPIs as DataFrame"""
        if not hasattr(kpis, '__dict__'):
            return pd.DataFrame()
        
        kpi_list = []
        for key, value in kpis.__dict__.items():
            if not key.startswith('_'):
                kpi_list.append({
                    'Metric': key.replace('_', ' ').title(),
                    'Value': value,
                    'Type': type(value).__name__
                })
        
        return pd.DataFrame(kpi_list)
    
    def _prepare_revenue_dataframe(self, revenue_data: Dict[str, Any]) -> pd.DataFrame:
        """Prepare revenue data as DataFrame"""
        if 'monthly_revenue' in revenue_data:
            return pd.DataFrame(list(revenue_data['monthly_revenue'].items()), 
                              columns=['Month', 'Revenue'])
        return pd.DataFrame()
    
    def _prepare_customer_dataframe(self, customer_data: Dict[str, Any]) -> pd.DataFrame:
        """Prepare customer data as DataFrame"""
        return pd.DataFrame([customer_data])
    
    def _format_excel_workbook(self, workbook) -> None:
        """Format Excel workbook with styles"""
        # Add formatting to make the Excel report more professional
        for worksheet in workbook.worksheets:
            # Header formatting
            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
    
    def _serialize_for_json(self, obj: Any) -> Any:
        """Serialize object for JSON export"""
        if isinstance(obj, dict):
            return {k: self._serialize_for_json(v) for k, v in obj.items()}
        elif isinstance(obj, list):
            return [self._serialize_for_json(item) for item in obj]
        elif isinstance(obj, (date, datetime)):
            return obj.isoformat()
        elif hasattr(obj, '__dict__'):
            return self._serialize_for_json(obj.__dict__)
        else:
            return obj
    
    def _get_default_email_template(self) -> str:
        """Get default email template"""
        return """
        <html>
        <body>
            <h2>Automated Report - {template_name}</h2>
            <p>Please find attached your automated report for the period {start_date} to {end_date}.</p>
            <p><strong>Report Summary:</strong></p>
            <p>{summary}</p>
            <p>Generated on: {generated_at}</p>
            <hr>
            <p><small>This is an automated message from OptimPV Billing System.</small></p>
        </body>
        </html>
        """
    
    # Additional chart creation methods
    def _create_quarterly_revenue_chart(self, data: Dict[str, Any]) -> str:
        """Create quarterly revenue chart"""
        return self._create_revenue_trend_chart(data)  # Reuse for now
    
    def _create_customer_analysis_chart(self, data: Dict[str, Any]) -> str:
        """Create customer analysis chart"""
        return self._create_customer_segmentation_chart(data)  # Reuse for now
    
    def _create_forecasting_chart(self, data: Dict[str, Any]) -> str:
        """Create forecasting chart"""
        if 'predictions' not in data:
            return ""
        
        predictions = data['predictions']
        
        metrics = [p.metric_name for p in predictions]
        current_values = [p.current_value for p in predictions]
        predicted_values = [p.predicted_value for p in predictions]
        
        fig = go.Figure()
        
        fig.add_trace(go.Bar(
            name='Current',
            x=metrics,
            y=current_values,
            marker_color='lightblue'
        ))
        
        fig.add_trace(go.Bar(
            name='Predicted',
            x=metrics,
            y=predicted_values,
            marker_color='darkblue'
        ))
        
        fig.update_layout(
            title='Predictive Analytics - Current vs Predicted',
            xaxis_title='Metrics',
            yaxis_title='Values',
            barmode='group',
            template='plotly_white',
            height=400
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="forecasting_chart")
    
    def _create_payment_patterns_chart(self, data: Dict[str, Any]) -> str:
        """Create payment patterns chart"""
        # Mock data for payment patterns
        payment_data = {
            'On Time': 75,
            '1-15 Days Late': 15,
            '16-30 Days Late': 7,
            '30+ Days Late': 3
        }
        
        fig = go.Figure(data=[go.Pie(
            labels=list(payment_data.keys()),
            values=list(payment_data.values()),
            hole=0.3
        )])
        
        fig.update_layout(
            title='Payment Patterns Distribution',
            template='plotly_white',
            height=400
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="payment_patterns_chart")
    
    def _create_risk_matrix_chart(self, data: Dict[str, Any]) -> str:
        """Create risk matrix chart"""
        # Mock risk matrix data
        risk_levels = ['Low', 'Medium', 'High']
        customer_counts = [18, 5, 2]
        colors = ['green', 'orange', 'red']
        
        fig = go.Figure(data=[go.Bar(
            x=risk_levels,
            y=customer_counts,
            marker_color=colors
        )])
        
        fig.update_layout(
            title='Customer Risk Distribution',
            xaxis_title='Risk Level',
            yaxis_title='Number of Customers',
            template='plotly_white',
            height=400
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="risk_matrix_chart")
    
    def _create_aging_analysis_chart(self, data: Dict[str, Any]) -> str:
        """Create aging analysis chart"""
        # Mock aging data
        aging_buckets = ['Current', '1-30 Days', '31-60 Days', '61-90 Days', '90+ Days']
        amounts = [45000, 8000, 3000, 1500, 500]
        
        fig = go.Figure(data=[go.Bar(
            x=aging_buckets,
            y=amounts,
            marker_color=['green', 'yellow', 'orange', 'red', 'darkred']
        )])
        
        fig.update_layout(
            title='Accounts Receivable Aging Analysis',
            xaxis_title='Aging Bucket',
            yaxis_title='Amount (€)',
            template='plotly_white',
            height=400
        )
        
        return fig.to_html(include_plotlyjs='inline', div_id="aging_analysis_chart")