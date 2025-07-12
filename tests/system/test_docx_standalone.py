"""
Test standalone du système DOCX (sans Streamlit)
===============================================

Test minimal pour vérifier la structure DOCX sans dépendances Streamlit.
"""

import sys
import os
from pathlib import Path

def test_docx_imports():
    """Test des imports python-docx"""
    print("🧪 Test des imports python-docx...")
    
    try:
        from docx import Document
        from docx.shared import Inches, Cm, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
        print("✅ python-docx importé avec succès")
        return True
    except ImportError as e:
        print(f"❌ python-docx non disponible: {e}")
        print("💡 Installez avec: pip install python-docx")
        return False

def test_basic_docx_creation():
    """Test de création de document DOCX basique"""
    print("\n🧪 Test de création DOCX basique...")
    
    try:
        from docx import Document
        from docx.shared import Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        # Créer un document
        doc = Document()
        print("✅ Document créé")
        
        # Ajouter du contenu
        title = doc.add_paragraph("Rapport Test OptimPV")
        title.style = 'Title'
        print("✅ Titre ajouté")
        
        # Ajouter un tableau
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Paramètre'
        hdr_cells[1].text = 'Valeur'
        
        # Ajouter des données au tableau
        test_data = [
            ("Puissance", "25.5 kWc"),
            ("Économie", "65 000 €"),
            ("TRI", "8.5 %")
        ]
        
        for param, value in test_data:
            row_cells = table.add_row().cells
            row_cells[0].text = param
            row_cells[1].text = value
        
        print("✅ Tableau ajouté")
        
        # Test de sauvegarde
        test_file = "test_output.docx"
        doc.save(test_file)
        
        # Vérifier la taille du fichier
        if os.path.exists(test_file):
            file_size = os.path.getsize(test_file)
            print(f"✅ Fichier sauvegardé: {test_file} ({file_size} bytes)")
            
            # Nettoyer
            os.remove(test_file)
            print("✅ Fichier test nettoyé")
        
        return True
        
    except Exception as e:
        print(f"❌ Erreur lors de la création DOCX: {e}")
        return False

def test_docx_structure():
    """Test de la structure des modules DOCX OptimPV"""
    print("\n🧪 Test de la structure des modules...")
    
    # Vérifier l'existence des fichiers
    base_path = Path("modules/reporting/docx_system")
    
    expected_files = [
        "__init__.py",
        "docx_generator.py", 
        "template_manager.py",
        "data_mapper.py",
        "chart_inserter.py"
    ]
    
    all_exist = True
    for file_name in expected_files:
        file_path = base_path / file_name
        if file_path.exists():
            print(f"✅ {file_name} trouvé")
        else:
            print(f"❌ {file_name} manquant")
            all_exist = False
    
    return all_exist

def test_placeholder_parsing():
    """Test du parsing des placeholders"""
    print("\n🧪 Test du parsing des placeholders...")
    
    try:
        import re
        
        # Pattern de placeholder comme dans DataMapper
        placeholder_pattern = re.compile(r'\{\{([A-Z_]+)\}\}')
        
        # Texte de test
        test_text = "Puissance: {{POWER_KWC}} kWc, Économie: {{TOTAL_SAVINGS}} €"
        
        # Extraire les placeholders
        placeholders = placeholder_pattern.findall(test_text)
        
        expected = ['POWER_KWC', 'TOTAL_SAVINGS']
        if placeholders == expected:
            print(f"✅ Placeholders extraits: {placeholders}")
            
            # Test de remplacement
            data_map = {
                'POWER_KWC': '25.5',
                'TOTAL_SAVINGS': '65,000'
            }
            
            result_text = test_text
            for placeholder in placeholders:
                if placeholder in data_map:
                    result_text = result_text.replace(f"{{{{{placeholder}}}}}", data_map[placeholder])
            
            expected_result = "Puissance: 25.5 kWc, Économie: 65,000 €"
            if result_text == expected_result:
                print(f"✅ Remplacement réussi: {result_text}")
                return True
            else:
                print(f"❌ Remplacement échoué: {result_text}")
                return False
        else:
            print(f"❌ Placeholders incorrects: {placeholders}")
            return False
            
    except Exception as e:
        print(f"❌ Erreur lors du test placeholder: {e}")
        return False

def test_template_structure():
    """Test de la structure des templates"""
    print("\n🧪 Test de la structure des templates...")
    
    try:
        # Vérifier l'existence du dossier templates
        templates_dir = Path("modules/reporting/templates")
        
        if not templates_dir.exists():
            print("⚠️ Dossier templates n'existe pas encore (sera créé au premier usage)")
            return True
        
        # Lister les fichiers DOCX
        docx_files = list(templates_dir.glob("*.docx"))
        print(f"✅ Dossier templates existe, {len(docx_files)} template(s) trouvé(s)")
        
        return True
        
    except Exception as e:
        print(f"❌ Erreur lors du test template structure: {e}")
        return False

def main():
    """Fonction principale des tests standalone"""
    print("🚀 Tests standalone du système DOCX OptimPV")
    print("=" * 50)
    
    tests = [
        test_docx_imports,
        test_docx_structure,
        test_placeholder_parsing,
        test_template_structure,
        test_basic_docx_creation
    ]
    
    results = []
    for test in tests:
        try:
            result = test()
            results.append(result)
        except Exception as e:
            print(f"❌ Erreur fatale dans {test.__name__}: {e}")
            results.append(False)
    
    print("\n" + "=" * 50)
    print("📊 Résumé des tests standalone:")
    
    total_tests = len(results)
    passed_tests = sum(results)
    failed_tests = total_tests - passed_tests
    
    print(f"✅ Tests réussis: {passed_tests}/{total_tests}")
    print(f"❌ Tests échoués: {failed_tests}/{total_tests}")
    
    if failed_tests == 0:
        print("\n🎉 Structure DOCX validée ! Système prêt à l'emploi.")
    elif passed_tests > 0:
        print(f"\n⚠️ Structure partiellement validée. Installez python-docx pour tester complètement.")
    else:
        print("\n❌ Problèmes de structure détectés.")
    
    return failed_tests == 0

if __name__ == "__main__":
    success = main()
    exit(0 if success else 1)